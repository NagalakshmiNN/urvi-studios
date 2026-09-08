"""
Builds the "What the tests check" workbook from Playwright's own list of
tests, so the spreadsheet can never drift from the suite itself.

Regenerate any time with:
    cd urvi-app
    PW_SKIP_SERVER=1 npx playwright test --list --reporter=json > /tmp/tests.json
    node scripts/flatten-tests.js          # writes /tmp/test-rows.json
    python scripts/build-test-inventory.py
"""

import json
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OLIVE = "3F4827"
GOLD = "A98238"
IVORY = "FBF8F2"
SAND = "F1EBE0"
LINE = "D8CFC0"

# One friendly area name per test file — the wording a shop owner would use,
# not the file name.
AREAS = {
    "unit/pricing-and-gst.spec.ts": ("Money & pricing rules", "Logic"),
    "unit/richtext.spec.ts": ("Product description formatting", "Logic"),
    "unit/whatsapp-and-status.spec.ts": ("WhatsApp links & order status wording", "Logic"),
    "api/checkout-order.spec.ts": ("Placing an order (behind the scenes)", "Endpoint"),
    "api/public-endpoints.spec.ts": ("Coupons, wishlist, contact form & access control", "Endpoint"),
    "api/razorpay-verify.spec.ts": ("Razorpay payment verification", "Endpoint"),
    "e2e/storefront.spec.ts": ("Homepage, style guides & information pages", "Browser"),
    "e2e/shop-filters.spec.ts": ("Shop page & filters", "Browser"),
    "e2e/product-and-cart.spec.ts": ("Product page & cart", "Browser"),
    "e2e/account.spec.ts": ("Customer accounts, wishlist & addresses", "Browser"),
    "e2e/checkout.spec.ts": ("Checkout, as a customer sees it", "Browser"),
    "e2e/admin-products.spec.ts": ("Admin — products", "Browser"),
    "e2e/admin-excel.spec.ts": ("Admin — Excel import & export", "Browser"),
    "e2e/admin-orders.spec.ts": ("Admin — orders & stock", "Browser"),
    "e2e/admin-dashboard.spec.ts": ("Admin — dashboard, coupons, messages & data export", "Browser"),
}

# Order the sheet the way someone would walk the site, not alphabetically.
AREA_ORDER = [
    "e2e/storefront.spec.ts",
    "e2e/shop-filters.spec.ts",
    "e2e/product-and-cart.spec.ts",
    "e2e/account.spec.ts",
    "e2e/checkout.spec.ts",
    "api/checkout-order.spec.ts",
    "api/razorpay-verify.spec.ts",
    "api/public-endpoints.spec.ts",
    "e2e/admin-products.spec.ts",
    "e2e/admin-excel.spec.ts",
    "e2e/admin-orders.spec.ts",
    "e2e/admin-dashboard.spec.ts",
    "unit/pricing-and-gst.spec.ts",
    "unit/richtext.spec.ts",
    "unit/whatsapp-and-status.spec.ts",
]

NOT_COVERED = [
    (
        "A real card payment going through",
        "No live Razorpay charge is ever made by the tests. What IS tested is the check that decides "
        "whether a payment is genuine — a correctly signed payment confirms the order and takes stock, "
        "a forged one changes nothing, and a repeated callback doesn't take stock twice.",
        "Place one small real test order yourself before relying on live payments.",
    ),
    (
        "Email actually arriving in an inbox",
        "Emails are captured to a file instead of being sent, so the tests can check who would have been "
        "emailed and what it said. That proves the site tries to send the right thing — not that Gmail "
        "delivered it.",
        "Check one real order end to end and confirm the email lands.",
    ),
    (
        "WhatsApp messages actually sending",
        "WhatsApp requires a person to tap Send — no system can send on your behalf without the paid "
        "WhatsApp Business API. The tests check the link and the pre-filled message are correct, which is "
        "everything the site controls.",
        "Tap one through on a real phone after a deploy.",
    ),
    (
        "How anything looks",
        "These tests check behaviour, not design. A page could pass every test and still look wrong, have "
        "a photo cropped badly, or read awkwardly on a phone.",
        "Your own eyes on a real phone, after each deploy.",
    ),
    (
        "Your live production data",
        "Every run builds its own separate throwaway database and deletes it afterwards. Real customers, "
        "orders and products are never touched.",
        "Nothing needed — this is deliberate.",
    ),
]


def build():
    rows = json.load(open("/tmp/test-rows.json"))
    by_file = {}
    for r in rows:
        by_file.setdefault(r["file"], []).append(r)

    wb = Workbook()

    # ---------------------------------------------------------------- Sheet 1
    ws = wb.active
    ws.title = "What The Tests Check"

    headers = ["#", "Area of the site", "Group", "What it checks", "Type", "Test file", "Line"]
    ws.append(headers)
    for i, _ in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(name="Arial", size=11, bold=True, color=IVORY)
        c.fill = PatternFill("solid", fgColor=OLIVE)
        c.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
    ws.row_dimensions[1].height = 30

    thin = Side(style="thin", color=LINE)
    border = Border(bottom=thin)

    n = 0
    banded = False
    last_area = None
    for path in AREA_ORDER:
        area, kind = AREAS[path]
        for r in by_file.get(path, []):
            n += 1
            if area != last_area:
                banded = not banded
                last_area = area
            ws.append([n, area, r["group"], r["title"], kind, path, r["line"]])
            for col in range(1, 8):
                c = ws.cell(row=n + 1, column=col)
                c.font = Font(name="Arial", size=10)
                c.alignment = Alignment(vertical="top", wrap_text=(col in (2, 3, 4)))
                c.border = border
                if banded:
                    c.fill = PatternFill("solid", fgColor=IVORY)
            ws.cell(row=n + 1, column=1).alignment = Alignment(horizontal="center", vertical="top")
            ws.cell(row=n + 1, column=5).alignment = Alignment(horizontal="center", vertical="top")
            ws.cell(row=n + 1, column=7).alignment = Alignment(horizontal="center", vertical="top")
            ws.cell(row=n + 1, column=6).font = Font(name="Arial", size=9, color="7A7268")

    widths = {"A": 5, "B": 30, "C": 26, "D": 68, "E": 11, "F": 30, "G": 6}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{n + 1}"

    # ---------------------------------------------------------------- Sheet 2
    s = wb.create_sheet("Summary", 0)
    s["A1"] = "Urvi Studios — automated test coverage"
    s["A1"].font = Font(name="Arial", size=16, bold=True, color=OLIVE)
    s["A2"] = f"Generated {date.today().strftime('%d %B %Y')} · every test in the suite, straight from the code"
    s["A2"].font = Font(name="Arial", size=10, italic=True, color="7A7268")

    s["A4"] = "These tests run by themselves every time a change is pushed to GitHub."
    s["A5"] = "A green tick next to the commit means everything below still works. A red cross means one of"
    s["A6"] = "these checks failed — open that run on GitHub and download the report to see which one."
    for row in (4, 5, 6):
        s[f"A{row}"].font = Font(name="Arial", size=10)

    s["A8"] = "Total tests"
    s["A8"].font = Font(name="Arial", size=11, bold=True)
    s["B8"] = f"=COUNTA('What The Tests Check'!A2:A{n + 1})"
    s["B8"].font = Font(name="Arial", size=11, bold=True, color=OLIVE)

    # By type
    s["A10"] = "By type"
    s["A10"].font = Font(name="Arial", size=12, bold=True, color=OLIVE)
    s["A11"], s["B11"], s["C11"] = "Type", "Tests", "What it means"
    type_notes = {
        "Browser": "A real browser clicks through the real pages, like a customer or you would",
        "Endpoint": "Calls the site's own machinery directly — orders, payments, security",
        "Logic": "The calculations on their own: GST, prices, phone numbers, wording",
    }
    for i, (t, note) in enumerate(type_notes.items(), start=12):
        s[f"A{i}"] = t
        s[f"B{i}"] = f"=COUNTIF('What The Tests Check'!$E$2:$E${n + 1},A{i})"
        s[f"C{i}"] = note

    # By area
    start = 16
    s[f"A{start - 1}"] = "By area of the site"
    s[f"A{start - 1}"].font = Font(name="Arial", size=12, bold=True, color=OLIVE)
    s[f"A{start}"], s[f"B{start}"] = "Area", "Tests"
    seen = []
    for path in AREA_ORDER:
        area = AREAS[path][0]
        if area not in seen:
            seen.append(area)
    for i, area in enumerate(seen, start=start + 1):
        s[f"A{i}"] = area
        s[f"B{i}"] = f"=COUNTIF('What The Tests Check'!$B$2:$B${n + 1},A{i})"
    total_row = start + len(seen) + 1
    s[f"A{total_row}"] = "Total"
    s[f"B{total_row}"] = f"=SUM(B{start + 1}:B{total_row - 1})"
    for cell in (f"A{total_row}", f"B{total_row}"):
        s[cell].font = Font(name="Arial", size=11, bold=True, color=OLIVE)

    for header_row in (11, start):
        for col in ("A", "B", "C"):
            c = s[f"{col}{header_row}"]
            c.font = Font(name="Arial", size=10, bold=True, color=IVORY)
            c.fill = PatternFill("solid", fgColor=OLIVE)

    for row in s.iter_rows(min_row=1, max_row=total_row, max_col=3):
        for c in row:
            if c.font is None or c.font.name != "Arial":
                c.font = Font(name="Arial", size=10)
    s.column_dimensions["A"].width = 46
    s.column_dimensions["B"].width = 10
    s.column_dimensions["C"].width = 72

    # ---------------------------------------------------------------- Sheet 3
    nc = wb.create_sheet("What Tests Can't Catch")
    nc["A1"] = "What these tests deliberately do not check"
    nc["A1"].font = Font(name="Arial", size=14, bold=True, color=OLIVE)
    nc["A2"] = "Worth knowing, so the green tick isn't mistaken for more than it is."
    nc["A2"].font = Font(name="Arial", size=10, italic=True, color="7A7268")

    nc.append([])
    nc.append(["Not covered", "Why", "What to do instead"])
    for col in ("A", "B", "C"):
        c = nc[f"{col}4"]
        c.font = Font(name="Arial", size=10, bold=True, color=IVORY)
        c.fill = PatternFill("solid", fgColor=GOLD)
        c.alignment = Alignment(vertical="center", wrap_text=True)

    for i, (what, why, instead) in enumerate(NOT_COVERED, start=5):
        nc[f"A{i}"], nc[f"B{i}"], nc[f"C{i}"] = what, why, instead
        for col in ("A", "B", "C"):
            c = nc[f"{col}{i}"]
            c.font = Font(name="Arial", size=10, bold=(col == "A"))
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = border
        nc.row_dimensions[i].height = 62

    nc.column_dimensions["A"].width = 34
    nc.column_dimensions["B"].width = 78
    nc.column_dimensions["C"].width = 46

    out = "/home/claude/deliverables/Urvi_Studios_Test_Coverage.xlsx"
    wb.save(out)
    print(f"wrote {out} with {n} tests")


if __name__ == "__main__":
    build()
